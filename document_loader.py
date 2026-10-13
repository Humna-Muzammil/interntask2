from pathlib import Path
import fitz
from docx import Document

def extract_file(path: str) -> list[dict]:
    """Return text segments with source and page/section metadata."""
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    segments = []

    if suffix == ".pdf":
        with fitz.open(file_path) as pdf:
            for page_num, page in enumerate(pdf, start=1):
                text = page.get_text("text").strip()
                if text:
                    segments.append({
                        "text": text,
                        "source": file_path.name,
                        "location": f"page {page_num}",
                    })
    elif suffix == ".txt":
        text = file_path.read_text(encoding="utf-8", errors="replace")
        if text.strip():
            segments.append({
                "text": text,
                "source": file_path.name,
                "location": "text file",
            })
    elif suffix == ".docx":
        doc = Document(file_path)
        current_section = "document"
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = (paragraph.style.name or "").lower() if paragraph.style else ""
            if "heading" in style_name:
                current_section = text
            segments.append({
                "text": text,
                "source": file_path.name,
                "location": f"section: {current_section}",
            })
        for table_index, table in enumerate(doc.tables, start=1):
            rows = []
            for row in table.rows:
                rows.append(" | ".join(cell.text.strip() for cell in row.cells))
            table_text = "\n".join(rows).strip()
            if table_text:
                segments.append({
                    "text": table_text,
                    "source": file_path.name,
                    "location": f"table {table_index}",
                })
    else:
        raise ValueError("Unsupported file type. Upload PDF, TXT, or DOCX.")

    if not segments:
        raise ValueError(f"No extractable text found in {file_path.name}. Scanned PDFs may need OCR.")
    return segments
